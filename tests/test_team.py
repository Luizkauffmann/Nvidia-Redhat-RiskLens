"""Phase 4 team: quant sandbox, Excel attribution pack, orchestrator delegation."""
from __future__ import annotations

import json
import os

import pytest

from src.agents.base import ExtraTool
from src.agents.team import SPECIALISTS, build_team
from src.llm.mock_provider import MockProvider
from src.llm.provider import LLMProvider, LLMResponse, ToolCall
from src.tools import ToolContext, call_tool, list_tools, set_context
from src.utils import config
from tests.conftest import CURR, PREV

PAIR = {"previous_date": PREV, "current_date": CURR}


@pytest.fixture(scope="module", autouse=True)
def ctx(generated, attribution, tmp_path_factory):
    c = ToolContext(data_dir=generated[0], cache_dir=tmp_path_factory.mktemp("cache"), disk_cache=False)
    c._mem[f"attr_default_CECL_{PREV}_{CURR}"] = attribution
    set_context(c)
    old = os.environ.get("RISKLENS_OUTPUT_DIR")
    os.environ["RISKLENS_OUTPUT_DIR"] = str(tmp_path_factory.mktemp("outputs"))
    config.reload()
    yield c
    set_context(None)
    if old is None:
        os.environ.pop("RISKLENS_OUTPUT_DIR", None)
    else:
        os.environ["RISKLENS_OUTPUT_DIR"] = old
    config.reload()


# --------------------------------------------------------------------------- quant sandbox
STATE_CODE = """
ecl = at(results[results.framework == "CECL"], CURR)
result = ecl.groupby("portfolio").ecl.sum().sort_values(ascending=False)
"""


def test_sandbox_runs_pandas_and_matches_engine_output(ctx):
    r = call_tool("run_python_analysis", {"code": STATE_CODE, "purpose": "ECL by portfolio"})
    assert r.status == "ok", r.error
    got = r.result["result"]["values"]
    booked = ctx.store.results_at(CURR, "CECL").groupby("portfolio").ecl.sum()
    for k, v in got.items():
        assert v == pytest.approx(booked[k])
    assert r.result["code_sha256"] and r.result["code"].strip() == STATE_CODE.strip()


@pytest.mark.parametrize("code,needle", [
    ("import socket; socket.socket()", "blocked"),
    ("open('/tmp/risklens_x', 'w').write('a')", "read-only"),
    ("import subprocess; subprocess.run(['ls'])", "blocked"),
    ("import os; os.remove('x')", "blocked"),
])
def test_sandbox_blocks_side_effects(code, needle):
    r = call_tool("run_python_analysis", {"code": code})
    assert r.status == "error" and needle in r.error["message"]


def test_sandbox_cannot_read_ground_truth(generated):
    gt = generated[1] / "ground_truth.json"
    r = call_tool("run_python_analysis", {"code": f"result = open({str(gt)!r}).read()"})
    assert r.status == "error" and "ground truth" in r.error["message"]


def test_sandbox_errors_are_recoverable():
    r = call_tool("run_python_analysis", {"code": "result = loans.no_such_column"})
    assert r.status == "error" and r.error["error_type"] == "code_error"
    assert "describe_datasets" in r.error["hint"]["tip"]
    r = call_tool("run_python_analysis", {"code": "x = 1"})
    assert "result" in r.error["message"]


def test_describe_datasets_lists_tables():
    r = call_tool("describe_datasets", {})
    tables = {t["table"]: t for t in r.result["tables"]}
    assert {"loans", "results", "macro", "model_metadata"} <= set(tables)
    assert "ecl" in tables["results"]["columns"]


# --------------------------------------------------------------------------- Excel pack
def test_attribution_workbook_ties_and_has_formulas():
    from openpyxl import load_workbook

    r = call_tool("build_attribution_workbook", PAIR)
    assert r.status == "ok", r.error
    wb = load_workbook(r.result["path"])
    assert {"Summary", "Waterfall", "Portfolio x driver", "Macro variables", "Top borrowers", "Sources"} <= set(wb.sheetnames)
    w = wb["Waterfall"]
    labels = [w.cell(i, 1).value for i in range(1, w.max_row + 1)]
    total_row = labels.index("Total explained") + 1
    assert w.cell(total_row, 2).value.startswith("=SUM(")
    impacts = [w.cell(i, 2).value for i in range(5, total_row)]
    assert sum(impacts) == pytest.approx(r.result["headline"]["change"], abs=0.01)
    assert wb["Summary"]["B6"].value == "=B5-B4"
    src = wb["Sources"]
    assert src.max_row - 1 == len(r.result["source_calls"]) >= 8


# --------------------------------------------------------------------------- team
def test_agents_config_has_new_specialists():
    assert set(SPECIALISTS) <= set(config.load_yaml("agents")["agents"])
    assert "run_python_analysis" in list_tools("quant_agent")
    assert list_tools("reporting_agent")[-1] == "build_attribution_workbook"


@pytest.mark.parametrize("question,specialist", [
    ("Which borrowers should I investigate first?", "credit_agent"),
    ("Does everything reconcile?", "reconciliation_agent"),
    ("Build the Excel attribution pack for the committee", "reporting_agent"),
])
def test_mock_team_delegates_and_stays_grounded(question, specialist):
    team = build_team(MockProvider())
    run = team.run(question)
    assert run.status == "answered"
    assert f"ask_{specialist}" in run.tools_used
    assert (run.grounding or {}).get("status") in ("PASS", "NO_NUMBERS")
    sub = next(e for e in run.events if e["type"] == "team")
    assert specialist in sub["specialists"]


def test_delegation_tool_reports_bad_arguments():
    tool = ExtraTool(build_team(MockProvider()).orchestrator.extra_tools["ask_credit_agent"].schema, lambda question: {})
    rec = tool.call({})
    assert rec.status == "error" and "question" in rec.error["message"]


class ScriptedProvider(LLMProvider):
    """Orchestrator asks the quant; the quant inspects tables, writes code, then quotes the result."""
    name, model = "scripted", "scripted"

    def chat(self, messages, tools=None, **kwargs) -> LLMResponse:
        system = messages[0]["content"]
        tool_msgs = [m for m in messages if m["role"] == "tool"]
        names = {t["function"]["name"] for t in tools or []}
        if "lead investigator" in system:
            if not tool_msgs:
                return LLMResponse(tool_calls=[ToolCall("o1", "ask_quant_agent", {"question": "CRE ECL on the current date?"})])
            finding = json.loads(tool_msgs[-1]["content"])["result"]
            return LLMResponse(content=f"Per the Quant Analyst: {finding['answer']}")
        if "Quant Analyst" in system:
            if not tool_msgs:
                return LLMResponse(tool_calls=[ToolCall("q1", "describe_datasets", {})])
            if len(tool_msgs) == 1:
                code = "c = at(results[results.framework == 'CECL'], CURR)\nresult = {'cre_ecl': c[c.portfolio == 'CRE'].ecl.sum()}"
                return LLMResponse(tool_calls=[ToolCall("q2", "run_python_analysis", {"code": code, "purpose": "CRE ECL"})])
            out = json.loads(tool_msgs[-1]["content"])["result"]["result"]
            assert "run_python_analysis" in names
            return LLMResponse(content=f"CRE CECL ECL on {CURR} is ${out['cre_ecl'] / 1e6:.2f}M [run_python_analysis].")
        raise AssertionError("unexpected agent")


def test_team_with_quant_code_is_grounded(ctx):
    run = build_team(ScriptedProvider()).run("How much CECL ECL does CRE carry now?")
    assert run.status == "answered" and run.tools_used == ["ask_quant_agent"]
    assert run.grounding["status"] == "PASS"
    booked = ctx.store.results_at(CURR, "CECL")
    expected = booked[booked.portfolio == "CRE"].ecl.sum() / 1e6
    assert f"${expected:.2f}M" in run.answer
