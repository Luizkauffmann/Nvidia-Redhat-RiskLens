"""Reporting Agent tools: the CECL attribution pack as an Excel workbook.

Finance and the ALLL committee review allowance movements in Excel, so the Reporting Agent hands
them one. The workbook is built by code from the same tools the agents call (never by the LLM):

* inputs (tool results) are blue, formulas are black, cross-sheet links are green;
* totals, changes, shares and the residual are live Excel formulas, so the pack re-ties if
  someone edits a driver, and a reviewer can see exactly how each figure is derived;
* a Sources sheet lists every tool call behind the numbers (call_id, parameters, result hash).
"""
from __future__ import annotations

from pathlib import Path
from typing import Any

from src.tools.base import Framework, ToolError, call_tool, get_context, tool
from src.utils import config

REPORT_AGENTS = ["orchestrator", "reporting_agent"]

BLUE, GREEN = "0000FF", "008000"
MONEY = '$#,##0;($#,##0);"-"'
PCT = '0.0%;(0.0%);"-"'


def _tool(name: str, args: dict[str, Any], calls: list[dict[str, Any]]) -> dict[str, Any]:
    rec = call_tool(name, args, caller="reporting_agent", log=None)
    if rec.status != "ok":
        raise ToolError(f"{name} failed while building the workbook: {rec.error['message']}", "dependency_error",
                        rec.error.get("hint"))
    calls.append({"call_id": rec.call_id, "tool_name": name, "parameters": rec.parameters,
                  "timestamp": rec.timestamp, "result_sha256": rec.result_sha256})
    return rec.result


@tool("reporting", REPORT_AGENTS, slow=True)
def build_attribution_workbook(previous_date: str, current_date: str, framework: Framework = "CECL",
                               top_n: int = 25) -> dict[str, Any]:
    """Write the allowance attribution pack as an Excel workbook (summary, waterfall, portfolio x driver, macro variables, top borrowers, models and weights, reconciliation, sources) and return its path and headline figures."""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
        from openpyxl.utils import get_column_letter
    except ImportError as exc:                         # pragma: no cover
        raise ToolError("openpyxl is not installed", "execution_error", {"fix": "pip install openpyxl"}) from exc

    pair = {"previous_date": previous_date, "current_date": current_date, "framework": framework}
    calls: list[dict[str, Any]] = []
    cmp_ = _tool("compare_allowance", pair, calls)
    wf = _tool("calculate_portfolio_attribution", pair, calls)
    seg = _tool("calculate_segment_attribution", pair, calls)
    mv = _tool("calculate_variable_attribution", pair, calls)
    top = _tool("identify_top_contributors", {**pair, "n": top_n}, calls)
    models = _tool("compare_model_versions", pair, calls)
    weights = _tool("compare_scenario_weights", pair, calls)
    overlays = _tool("compare_management_overlays", pair, calls)
    rec = _tool("reconcile_allowance_change", pair, calls)

    wb = Workbook()
    font = "Arial"
    base, bold = Font(name=font, size=10), Font(name=font, size=10, bold=True)
    title_f = Font(name=font, size=14, bold=True)
    inp, link = Font(name=font, size=10, color=BLUE), Font(name=font, size=10, color=GREEN)
    head_fill = PatternFill("solid", fgColor="D9E1F2")
    thin = Side(style="thin", color="999999")

    def header(ws, row: int, labels: list[str]) -> None:
        for j, lab in enumerate(labels, start=1):
            c = ws.cell(row=row, column=j, value=lab)
            c.font, c.fill = bold, head_fill
            c.alignment = Alignment(horizontal="center", wrap_text=True)
            c.border = Border(bottom=thin)

    def put(ws, row: int, col: int, value: Any, kind: str = "text", fmt: str | None = None):
        c = ws.cell(row=row, column=col, value=value)
        c.font = {"input": inp, "link": link, "bold": bold}.get(kind, base)
        if fmt:
            c.number_format = fmt
        return c

    def widths(ws, cols: dict[str, int]) -> None:
        for col, w in cols.items():
            ws.column_dimensions[col].width = w

    # ---------------------------------------------------------------- Summary
    ws = wb.active
    ws.title = "Summary"
    put(ws, 1, 1, f"{framework} allowance attribution: {previous_date} to {current_date}").font = title_f
    put(ws, 2, 1, "Built by RiskLens from deterministic tools. Blue = tool output, black = formula, green = link.")
    rows = [("Previous allowance", cmp_["previous_allowance"], "input", MONEY),
            ("Current allowance", cmp_["current_allowance"], "input", MONEY),
            ("Change", "=B5-B4", "formula", MONEY),
            ("Change %", "=IFERROR(B6/B4,0)", "formula", PCT),
            ("Explained by attribution", "=Waterfall!B{explained}", "link", MONEY),
            ("Residual (change - explained)", "=B6-B8", "formula", MONEY),
            ("Previous balance", cmp_["previous_balance"], "input", MONEY),
            ("Current balance", cmp_["current_balance"], "input", MONEY),
            ("Reconciliation status (tool)", rec["overall_status"], "input", None),
            ("Primary driver (tool)", next(d["label"] for d in wf["drivers"] if d["driver"] == wf["primary_driver"]), "input", None),
            ("Any model version changed (tool)", "Yes" if models["any_model_version_changed"] else "No", "input", None),
            ("Scenario weights changed (tool)", "Yes" if weights["weights_changed"] else "No", "input", None)]
    header(ws, 3, ["Item", "Value"])
    for i, (lab, val, kind, fmt) in enumerate(rows, start=4):
        put(ws, i, 1, lab)
        put(ws, i, 2, val, kind, fmt)
    widths(ws, {"A": 36, "B": 24})

    # ---------------------------------------------------------------- Waterfall
    w = wb.create_sheet("Waterfall")
    put(w, 1, 1, "Attribution waterfall (sequential counterfactual; residual measured, never forced)").font = bold
    header(w, 3, ["Driver", "Impact ($)", "Share of change"])
    put(w, 4, 1, f"Allowance {previous_date}")
    put(w, 4, 2, "=Summary!B4", "link", MONEY)
    r = 5
    first = r
    for d in wf["drivers"]:
        put(w, r, 1, d["label"])
        put(w, r, 2, d["impact"], "input", MONEY)
        put(w, r, 3, f"=IFERROR(B{r}/Summary!$B$6,0)", "formula", PCT)
        r += 1
    last = r - 1
    explained = r
    put(w, r, 1, "Total explained", "bold")
    put(w, r, 2, f"=SUM(B{first}:B{last})", "formula", MONEY).font = bold
    put(w, r, 3, f"=IFERROR(B{r}/Summary!$B$6,0)", "formula", PCT)
    put(w, r + 1, 1, f"Allowance {current_date} (walk)")
    put(w, r + 1, 2, f"=B4+B{explained}", "formula", MONEY)
    put(w, r + 2, 1, "Check: walk minus reported current")
    put(w, r + 2, 2, f"=B{r + 1}-Summary!B5", "formula", MONEY)
    ws["B8"] = f"=Waterfall!B{explained}"
    ws["B8"].font, ws["B8"].number_format = link, MONEY
    widths(w, {"A": 40, "B": 18, "C": 16})

    # ---------------------------------------------------------------- Portfolio x driver
    s = wb.create_sheet("Portfolio x driver")
    keys = [d["driver"] for d in wf["drivers"]]
    labels = seg.get("driver_labels") or {d["driver"]: d["label"] for d in wf["drivers"]}
    header(s, 1, ["Portfolio", *[labels.get(k, k) for k in keys], "Total"])
    segs = [x for x in seg["segments"] if "segment" in x]
    for i, x in enumerate(segs, start=2):
        put(s, i, 1, x["segment"])
        for j, k in enumerate(keys, start=2):
            put(s, i, j, x["drivers"].get(k, 0.0), "input", MONEY)
        lc = get_column_letter(len(keys) + 1)
        put(s, i, len(keys) + 2, f"=SUM(B{i}:{lc}{i})", "formula", MONEY)
    tot = len(segs) + 2
    put(s, tot, 1, "Total", "bold")
    for j in range(2, len(keys) + 3):
        col = get_column_letter(j)
        put(s, tot, j, f"=SUM({col}2:{col}{tot - 1})", "formula", MONEY).font = bold
    widths(s, {"A": 16, **{get_column_letter(j): 18 for j in range(2, len(keys) + 3)}})

    # ---------------------------------------------------------------- Macro variables
    m = wb.create_sheet("Macro variables")
    header(m, 1, ["Variable", "Shapley impact ($)", "Sequential impact ($)", "Share of macro"])
    vs = mv["variables"]
    for i, v in enumerate(vs, start=2):
        put(m, i, 1, v["variable"])
        put(m, i, 2, v["impact"], "input", MONEY)
        put(m, i, 3, v.get("sequential_impact"), "input", MONEY)
        put(m, i, 4, f"=IFERROR(B{i}/$B${len(vs) + 2},0)", "formula", PCT)
    put(m, len(vs) + 2, 1, "Total macro effect", "bold")
    put(m, len(vs) + 2, 2, f"=SUM(B2:B{len(vs) + 1})", "formula", MONEY)
    put(m, len(vs) + 2, 3, f"=SUM(C2:C{len(vs) + 1})", "formula", MONEY)
    widths(m, {"A": 22, "B": 20, "C": 22, "D": 16})

    # ---------------------------------------------------------------- Top borrowers
    t = wb.create_sheet("Top borrowers")
    cols = ["rank", "counterparty_id", "customer_name", "portfolio", "balance_previous", "balance_current",
            "ecl_previous", "ecl_current", "ecl_change", "pd_12m_previous", "pd_12m_current", "main_driver", "reason_codes"]
    header(t, 1, [c.replace("_", " ") for c in cols])
    contributors = [c for c in top["contributors"] if "rank" in c]
    for i, c in enumerate(contributors, start=2):
        for j, col in enumerate(cols, start=1):
            val = c.get(col)
            if col == "reason_codes":
                val = "; ".join(val or [])
            if col == "ecl_change":
                put(t, i, j, f"=H{i}-G{i}", "formula", MONEY)
                continue
            fmt = MONEY if col.startswith(("balance", "ecl")) else ("0.00%" if col.startswith("pd_") else None)
            put(t, i, j, val, "input" if fmt else "text", fmt)
    widths(t, {"A": 6, "B": 14, "C": 28, "D": 10, "E": 16, "F": 16, "G": 14, "H": 14, "I": 14, "J": 10, "K": 10,
               "L": 22, "M": 70})

    # ---------------------------------------------------------------- Models, weights, overlays
    g = wb.create_sheet("Models and weights")
    put(g, 1, 1, "Scenario weights", "bold")
    header(g, 2, ["Scenario", f"Weight {previous_date}", f"Weight {current_date}", "Change (pp)"])
    sc = weights["scenarios"]
    for i, x in enumerate(sc, start=3):
        put(g, i, 1, x["scenario"])
        put(g, i, 2, x["weight_previous"], "input", PCT)
        put(g, i, 3, x["weight_current"], "input", PCT)
        put(g, i, 4, f"=(C{i}-B{i})*100", "formula", '0.0;(0.0);"-"')
    put(g, len(sc) + 3, 1, "Total")
    put(g, len(sc) + 3, 2, f"=SUM(B3:B{len(sc) + 2})", "formula", PCT)
    put(g, len(sc) + 3, 3, f"=SUM(C3:C{len(sc) + 2})", "formula", PCT)
    r = len(sc) + 5
    put(g, r, 1, "Allowance impact of the weight change (tool)")
    put(g, r, 2, weights["allowance_impact_of_weight_change"], "input", MONEY)
    put(g, r + 1, 1, "Management overlays: previous / current (tool)")
    put(g, r + 1, 2, overlays["overlay_total_previous"], "input", MONEY)
    put(g, r + 1, 3, overlays["overlay_total_current"], "input", MONEY)
    r += 3
    put(g, r, 1, "Model versions by product", "bold")
    header(g, r + 1, ["Product", "PD previous", "PD current", "LGD previous", "LGD current", "EAD previous", "EAD current"])
    for i, x in enumerate(models["by_product"], start=r + 2):
        vals = [x["product_type"], x.get("pd_model_previous"), x.get("pd_model_current"), x.get("lgd_model_previous"),
                x.get("lgd_model_current"), x.get("ead_model_previous"), x.get("ead_model_current")]
        for j, v in enumerate(vals, start=1):
            put(g, i, j, v, "input" if j > 1 else "text")
    widths(g, {"A": 44, "B": 18, "C": 18, "D": 16, "E": 16, "F": 16, "G": 16})

    # ---------------------------------------------------------------- Reconciliation
    q = wb.create_sheet("Reconciliation")
    header(q, 1, ["Control", "Status", "Value"])
    for i, c in enumerate(rec.get("controls", []), start=2):
        put(q, i, 1, c.get("control"))
        put(q, i, 2, c.get("status"), "input")
        v = c.get("value")
        put(q, i, 3, v if isinstance(v, (int, float)) else str(v), "input")
    widths(q, {"A": 48, "B": 12, "C": 24})

    # ---------------------------------------------------------------- Sources
    src = wb.create_sheet("Sources")
    header(src, 1, ["call_id", "tool", "parameters", "timestamp (UTC)", "result sha256"])
    for i, c in enumerate(calls, start=2):
        for j, v in enumerate([c["call_id"], c["tool_name"], str(c["parameters"]), c["timestamp"], c["result_sha256"]], 1):
            put(src, i, j, v)
    widths(src, {"A": 16, "B": 32, "C": 70, "D": 28, "E": 20})

    wb.calculation.fullCalcOnLoad = True                  # Excel / LibreOffice compute the formulas on open
    out_dir = config.path("outputs") / "reports"
    out_dir.mkdir(parents=True, exist_ok=True)
    path = out_dir / f"{framework.lower()}_attribution_{previous_date}_{current_date}.xlsx"
    wb.save(path)
    return {"path": str(path), "framework": framework, "previous_date": previous_date, "current_date": current_date,
            "sheets": wb.sheetnames,
            "headline": {"previous_allowance": cmp_["previous_allowance"], "current_allowance": cmp_["current_allowance"],
                         "change": cmp_["change"], "explained_change": wf["explained_change"], "residual": wf["residual"],
                         "primary_driver": wf["primary_driver"], "reconciliation": rec["overall_status"],
                         "top_borrower": contributors[0]["counterparty_id"] if contributors else None},
            "source_calls": calls,
            "notes": "Totals, changes, shares and the residual are Excel formulas; inputs come from the listed tool calls."}
